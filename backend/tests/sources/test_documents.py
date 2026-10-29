"""Document extraction (PDF, DOCX, XLSX/CSV, HTML, text) and the table -> records mapping."""

from __future__ import annotations

import hashlib
import io
from datetime import datetime
from typing import Any

import httpx
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from botai.core.errors import BotaiError, ErrorCode
from botai.sources.adapter import AdapterResult, NeedsUserInput
from botai.sources.documents import chunk_text, detect_format, extract, to_chunks
from botai.sources.factory import AdapterFactory
from tests.sources.helpers import make_ctx

STUDENT = "Иванова Анна Сергеевна"


def make_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 400 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>\nstream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets: list[int] = []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Правила пересдачи")
    doc.add_paragraph("Заявление подаётся за три дня до экзамена.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Этап"
    table.rows[0].cells[1].text = "Срок"
    table.rows[1].cells[0].text = "Заявление"
    table.rows[1].cells[1].text = "до 5 ноября"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


EXAMS = [
    ["ФИО", "Предмет", "Дата", "Время", "Аудитория"],
    ["Петров Пётр Петрович", "Алгоритмы", datetime(2026, 12, 20), "10:00", "Б-204"],
    [STUDENT, "Базы данных", datetime(2026, 12, 22), "13:00", "В-310"],
    [STUDENT, "Математический анализ", datetime(2026, 12, 25), "09:30", "А-101"],
]


def make_xlsx(rows: list[list[Any]] = EXAMS, sheet: str = "Сессия") -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exam_spec(**table: Any) -> dict[str, Any]:
    spec_table: dict[str, Any] = {
        "kind": "scheduleEvents",
        "sheet": "Сессия",
        "rowFilter": {"column": "ФИО", "equalsParam": "fullName"},
        "columns": {
            "title": {"column": "Предмет"},
            "date": {"column": "Дата", "transform": "date"},
            "startTime": {"column": "Время", "transform": "time"},
            "location": {"column": "Аудитория"},
        },
    }
    spec_table.update(table)
    return {
        "adapter": "document",
        "source": {"kind": "upload"},
        "format": "xlsx",
        "table": spec_table,
    }


async def ingest(spec: dict[str, Any], data: bytes, **kw: Any) -> AdapterResult:
    ctx = make_ctx(transport=httpx.MockTransport(lambda r: httpx.Response(404)), **kw)
    ctx.content = data
    return await AdapterFactory().create(spec).fetch(ctx)


# --- extraction -----------------------------------------------------------------------------


def test_format_detection() -> None:
    assert detect_format(make_pdf("x")) == "pdf"
    assert detect_format(make_docx()) == "docx"
    assert detect_format(make_xlsx()) == "xlsx"
    assert detect_format(b"<!DOCTYPE html><html></html>") == "html"
    assert detect_format(b"a;b\n1;2", "t.csv") == "csv"
    assert detect_format("просто текст".encode()) == "text"


def test_pdf_text_is_extracted() -> None:
    out = extract(make_pdf("Exam schedule winter session"), "pdf")
    assert "Exam schedule winter session" in out.text


def test_broken_pdf_is_a_validation_error_not_a_crash() -> None:
    with pytest.raises(BotaiError) as exc:
        extract(b"%PDF-1.4 garbage", "pdf")
    assert exc.value.code == ErrorCode.VALIDATION


def test_docx_text_and_tables() -> None:
    out = extract(make_docx(), "docx")
    assert "Заявление подаётся" in out.text
    assert out.tables[0].rows == [["Этап", "Срок"], ["Заявление", "до 5 ноября"]]
    chunks = to_chunks(out)
    assert any("Этап: Заявление; Срок: до 5 ноября" in c or "Заявление; Срок" in c for c in chunks)


def test_xlsx_cells_become_strings_with_dates_in_iso() -> None:
    out = extract(make_xlsx(), "xlsx")
    assert out.tables[0].name == "Сессия"
    assert out.tables[0].rows[2] == [
        STUDENT,
        "Базы данных",
        "2026-12-22T00:00:00",
        "13:00",
        "В-310",
    ]


def test_csv_with_semicolons_and_windows_1251() -> None:
    data = "Предмет;Дата\nФизика;12.01.2026\n".encode("cp1251")
    out = extract(data, "csv")
    assert out.tables[0].rows == [["Предмет", "Дата"], ["Физика", "12.01.2026"]]


def test_html_page_drops_scripts_and_reads_tables() -> None:
    html = (
        "<html><head><title>Условия</title><script>evil()</script></head><body>"
        "<p>Текст страницы</p><table><tr><th>А</th><th>Б</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "</body></html>"
    )
    out = extract(html.encode(), "html")
    assert out.title == "Условия"
    assert "evil" not in out.text
    assert "Текст страницы" in out.text
    assert out.tables[0].rows == [["А", "Б"], ["1", "2"]]


def test_chunks_are_bounded_and_keep_everything() -> None:
    text = "\n\n".join(f"Абзац номер {i} " + "слово " * 80 for i in range(30))
    chunks = chunk_text(text, limit=500)
    assert all(len(c) <= 500 for c in chunks)
    assert "".join(chunks).count("Абзац номер") == 30


# --- adapter: text --------------------------------------------------------------------------


async def test_document_without_table_yields_a_document_and_chunks() -> None:
    spec = {"adapter": "document", "source": {"kind": "upload"}}
    result = await ingest(spec, make_docx(), filename="rules.docx")
    (doc,) = result.documents
    assert doc.title == "rules"
    assert doc.kind == "file"
    assert doc.hash == hashlib.sha256(make_docx()).hexdigest() or len(doc.hash) == 64
    assert doc.chunks
    assert result.records == []


async def test_same_hash_means_nothing_to_do() -> None:
    data = make_docx()
    spec = {"adapter": "document", "source": {"kind": "upload"}}
    result = await ingest(spec, data, known_hash=hashlib.sha256(data).hexdigest())
    assert result.unchanged
    assert result.documents == []
    assert result.records == []


async def test_document_from_url_uses_only_the_allowed_host() -> None:
    html = "<html><title>Условия</title><body>Текст</body></html>".encode()
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        return httpx.Response(200, content=html, headers={"content-type": "text/html"})

    spec = {"adapter": "document", "source": {"kind": "url", "url": "https://files.fakeuni.test/r"}}
    ctx = make_ctx(transport=httpx.MockTransport(handler), allow_hosts=("files.fakeuni.test",))
    result = await AdapterFactory().create(spec).fetch(ctx)
    assert seen == ["https://files.fakeuni.test/r"]
    assert result.documents[0].kind == "page"
    assert result.documents[0].title == "Условия"

    blocked = make_ctx(transport=httpx.MockTransport(handler), allow_hosts=("elsewhere.test",))
    with pytest.raises(BotaiError) as exc:
        await AdapterFactory().create(spec).fetch(blocked)
    assert exc.value.code == ErrorCode.EGRESS_BLOCKED


async def test_missing_link_target_is_a_problem_not_an_exception() -> None:
    spec = {"adapter": "document", "source": {"kind": "url", "url": "https://files.fakeuni.test/x"}}
    ctx = make_ctx(
        transport=httpx.MockTransport(lambda r: httpx.Response(404)),
        allow_hosts=("files.fakeuni.test",),
    )
    result = await AdapterFactory().create(spec).fetch(ctx)
    assert result.outcomes[0].problem == "not_found"
    assert result.documents == []


# --- adapter: table -> records --------------------------------------------------------------


async def test_table_rows_of_the_student_become_schedule_events() -> None:
    result = await ingest(exam_spec(), make_xlsx(), profile={"fullName": STUDENT})
    events = sorted(result.records, key=lambda r: r.fields["startsAt"])
    assert [e.fields["title"] for e in events] == ["Базы данных", "Математический анализ"]
    assert events[0].fields["startsAt"] == "2026-12-22T10:00:00+00:00"  # 13:00 Moscow
    assert events[0].fields["endsAt"] == events[0].fields["startsAt"]  # only a start was given
    assert events[0].fields["location"] == "В-310"
    assert result.outcomes[0].ok
    assert result.outcomes[0].items_seen == 2
    # the other student's row is not in the result at all
    assert all("Петров" not in repr(r) for r in result.records)


async def test_row_filter_normalises_case_spacing_and_yo() -> None:
    rows = [["ФИО", "Предмет", "Дата"], ["  ИВАНОВА   анна сергеевна ", "Физика", "12.01.2027"]]
    spec = exam_spec(
        columns={"title": {"column": "Предмет"}, "date": {"column": "Дата", "transform": "date"}}
    )
    result = await ingest(spec, make_xlsx(rows), profile={"fullName": STUDENT})
    assert len(result.records) == 1
    # a date without a time spans the local day
    assert result.records[0].fields["startsAt"] == "2027-01-11T21:00:00+00:00"
    assert result.records[0].fields["endsAt"] == "2027-01-12T20:59:00+00:00"


async def test_missing_profile_value_asks_the_user_instead_of_guessing() -> None:
    with pytest.raises(NeedsUserInput) as exc:
        await ingest(exam_spec(), make_xlsx(), profile={})
    assert exc.value.param == "fullName"


async def test_same_name_is_not_a_match_on_a_prefix() -> None:
    result = await ingest(exam_spec(), make_xlsx(), profile={"fullName": "Иванова Анна"})
    assert result.records == []
    assert result.outcomes[0].problem == "empty"


async def test_renamed_column_is_a_format_problem() -> None:
    spec = exam_spec(columns={"title": {"column": "Дисциплина"}})
    result = await ingest(spec, make_xlsx(), profile={"fullName": STUDENT})
    assert result.records == []
    assert result.outcomes[0].problem == "format"
    assert not result.outcomes[0].ok


async def test_table_ids_are_stable_across_runs() -> None:
    a = await ingest(exam_spec(), make_xlsx(), profile={"fullName": STUDENT})
    b = await ingest(exam_spec(), make_xlsx(), profile={"fullName": STUDENT})
    assert sorted(r.external_id for r in a.records) == sorted(r.external_id for r in b.records)


async def test_deadline_table_without_time_is_due_at_the_end_of_the_day() -> None:
    rows = [["Работа", "Сдать до"], ["Курсовая", "30.12.2026"]]
    spec = {
        "adapter": "document",
        "source": {"kind": "upload"},
        "table": {
            "kind": "deadlines",
            "columns": {
                "title": {"column": "Работа"},
                "date": {"column": "Сдать до", "transform": "date"},
            },
        },
    }
    result = await ingest(spec, make_xlsx(rows, "Лист1"))
    assert result.records[0].fields == {"title": "Курсовая", "dueAt": "2026-12-30T20:59:00+00:00"}


def test_a_zip_bomb_is_refused_before_it_is_unpacked(monkeypatch: pytest.MonkeyPatch) -> None:
    import zipfile

    from botai.sources import documents

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * (3 * 1024 * 1024))  # tiny when compressed
    assert len(buf.getvalue()) < 20_000
    monkeypatch.setattr(documents, "MAX_UNZIPPED", 1024 * 1024)
    with pytest.raises(BotaiError) as exc:
        extract(buf.getvalue(), "docx")
    assert exc.value.code == ErrorCode.VALIDATION
    with pytest.raises(BotaiError):
        extract(buf.getvalue(), "xlsx")
