"""Owner: W3. HTML-only pages (for HTML recipes)."""

from __future__ import annotations

from datetime import UTC, datetime
from html import escape

from fastapi import APIRouter, Request
from fastapi.responses import HTMLResponse, Response

from tests.fakeuni import data
from tests.fakeuni.login import is_logged_in

router = APIRouter()

NEWS = [
    ("/news/1", "Перенос занятий 4 ноября", "2026-10-27", "27 октября"),
    ("/news/2", "Запись на пересдачу открыта", "2026-10-25", "25 октября"),
]


def _state() -> object:
    from tests.fakeuni import api

    return api.STATE


@router.get("/news", response_class=HTMLResponse)
async def news() -> str:
    markup = getattr(_state(), "news_markup", "v1")
    if markup == "v2":
        items = "\n".join(
            f'<article class="post"><h3><a href="{href}">{escape(title)}</a></h3>'
            f'<span class="when" data-iso="{iso}">{human}</span></article>'
            for href, title, iso, human in NEWS
        )
        return f"<!doctype html><html><body><main><section>{items}</section></main></body></html>"
    items = "\n".join(
        f'<li class="news-item"><a href="{href}">{escape(title)}</a>'
        f'<time datetime="{iso}">{human}</time></li>'
        for href, title, iso, human in NEWS
    )
    return (
        '<!doctype html><html><body><main><ul class="news-list">\n'
        f"{items}\n</ul></main></body></html>"
    )


@router.get("/portal/deadlines", response_model=None)
async def deadlines_page(request: Request) -> Response:
    """Deadlines as an HTML table behind the login (a session that died redirects to sign-in)."""
    if not is_logged_in(request):
        return HTMLResponse(status_code=302, headers={"location": "/auth/login"})
    rows = "\n".join(
        f'<tr data-task="{d["taskId"]}"><td class="name">{escape(str(d["name"]))}</td>'
        f'<td class="course">{escape(str(d["course"]))}</td>'
        f'<td class="due" data-ts="{d["due"]}">'
        f"{datetime.fromtimestamp(int(str(d['due'])), tz=UTC):%d.%m.%Y %H:%M}</td></tr>"
        for d in data.deadlines()
    )
    return HTMLResponse(
        "<!doctype html><html><head><title>Задания</title></head><body>"
        f'<table id="tasks"><thead><tr><th>Задание</th><th>Курс</th><th>Срок</th></tr></thead>'
        f"<tbody>{rows}</tbody></table></body></html>"
    )


# --- files (for document sources) --------------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
STUDENT = "Иванова Анна Сергеевна"


def exam_rows(version: str = "v1") -> list[list[str]]:
    moved = version == "v2"
    return [
        ["ФИО", "Предмет", "Дата", "Время", "Аудитория"],
        ["Петров Пётр Петрович", "Алгоритмы", "20.12.2026", "10:00", "Б-204"],
        [STUDENT, "Базы данных", "22.12.2026", "13:00", "В-310"],
        [
            STUDENT,
            "Математический анализ",
            "26.12.2026" if moved else "25.12.2026",
            "09:30",
            "Г-1" if moved else "А-101",
        ],
    ]


def exams_xlsx(version: str = "v1") -> bytes:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Сессия"
    for row in exam_rows(version):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rules_docx() -> bytes:
    import io

    import docx

    doc = docx.Document()
    doc.add_paragraph("Правила пересдачи")
    doc.add_paragraph("Заявление на пересдачу подаётся за три дня до экзамена.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("/files/exams.xlsx", response_model=None)
async def exams_file() -> Response:
    version = getattr(_state(), "exams_version", "v1")
    return Response(exams_xlsx(version), media_type=XLSX)


@router.get("/files/rules.docx", response_model=None)
async def rules_file() -> Response:
    return Response(rules_docx(), media_type=DOCX)
