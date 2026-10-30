"""Turning files into text and tables: PDF, DOCX, XLSX/CSV, HTML, plain text.

Deterministic parsing only; a model is involved just once per changed document, when the agent
proposes which table columns mean what (`learn`), never here. Output is bounded: a hostile or
huge file cannot exhaust memory.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Literal

from botai.core.errors import BotaiError, ErrorCode
from botai.sources.extract import clean_text, collapse, node_text, parse_html

DocFormat = Literal["pdf", "docx", "xlsx", "csv", "html", "text"]
MAX_PDF_PAGES = 300
MAX_ROWS = 5000
MAX_COLS = 40
MAX_TEXT_CHARS = 2_000_000
MAX_UNZIPPED = 120 * 1024 * 1024  # DOCX/XLSX are zip files: a small file may unpack to gigabytes
MAX_ZIP_ENTRIES = 5000
CHUNK_CHARS = 1200
MAX_UPLOAD = 20 * 1024 * 1024  # one uploaded file: a chat attachment or a document
UPLOAD_ORIGIN = "https://upload.botai.local"  # uploads have no host; the allowlist stays empty


@dataclass(slots=True)
class Table:
    name: str
    rows: list[list[str]]


@dataclass(slots=True)
class Extracted:
    format: DocFormat
    text: str
    tables: list[Table] = field(default_factory=list)
    title: str | None = None


def detect_format(
    data: bytes, filename: str | None = None, content_type: str | None = None
) -> DocFormat:
    head = data[:2048].lstrip()
    lower_name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = zf.namelist()
        except zipfile.BadZipFile:
            names = []
        if any(n.startswith("word/") for n in names):
            return "docx"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        raise BotaiError(ErrorCode.VALIDATION, "Не удалось определить тип документа")
    utf16 = data.startswith((b"\xff\xfe", b"\xfe\xff"))  # byte order mark: text full of NULs
    low = head[:200].lower()
    if (
        low.startswith((b"<!doctype html", b"<html"))
        or "html" in ctype
        or lower_name.endswith((".html", ".htm"))
    ):
        return "html"
    if lower_name.endswith((".csv", ".tsv")) or "csv" in ctype:
        return "csv"
    if lower_name.endswith(".pdf") or "pdf" in ctype:
        raise BotaiError(ErrorCode.VALIDATION, "Файл не похож на PDF")
    if b"\x00" in data[:4096] and not utf16:
        raise BotaiError(ErrorCode.VALIDATION, "Этот тип файла не поддерживается")
    return "text"


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")  # the byte order mark picks the byte order
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(collapse(str(value)), 2000)


def _trim_rows(rows: list[list[str]]) -> list[list[str]]:
    out: list[list[str]] = []
    for row in rows:
        cells = row[:MAX_COLS]
        while cells and not cells[-1].strip():
            cells.pop()  # sheets are padded with empty cells to the used range
        if cells:
            out.append(cells)
    return out[:MAX_ROWS]


# --- per format -----------------------------------------------------------------------------


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise BotaiError(ErrorCode.VALIDATION, "PDF защищён паролем, откройте его без пароля")
        pages: list[str] = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")
    except PyPdfError as exc:
        raise BotaiError(ErrorCode.VALIDATION, "PDF не удалось прочитать") from exc
    text = clean_text("\n\n".join(p.strip() for p in pages if p.strip()), MAX_TEXT_CHARS)
    return Extracted("pdf", text)


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            entries = zf.infolist()
    except zipfile.BadZipFile:
        raise BotaiError(ErrorCode.VALIDATION, "Файл повреждён") from None
    if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNZIPPED:
        raise BotaiError(ErrorCode.VALIDATION, "Файл слишком велик после распаковки")


def _docx(data: bytes) -> Extracted:
    import docx

    _check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise BotaiError(ErrorCode.VALIDATION, "DOCX не удалось прочитать") from exc
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    tables: list[Table] = []
    for i, tbl in enumerate(document.tables):
        rows = [[_cell(c.text) for c in row.cells] for row in tbl.rows]
        tables.append(Table(f"table{i + 1}", _trim_rows(rows)))
    text = clean_text("\n".join(paragraphs), MAX_TEXT_CHARS)
    return Extracted("docx", text, tables)


def _xlsx(data: bytes) -> Extracted:
    from openpyxl import load_workbook

    _check_zip(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise BotaiError(ErrorCode.VALIDATION, "Таблицу XLSX не удалось прочитать") from exc
    tables: list[Table] = []
    try:
        for ws in workbook.worksheets:
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS * 2, max_col=MAX_COLS):
                rows.append([_cell(c) for c in row])
            tables.append(Table(ws.title, _trim_rows(rows)))
    finally:
        workbook.close()
    return Extracted("xlsx", "", tables)


def _csv(data: bytes) -> Extracted:
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    return Extracted("csv", "", [Table("csv", _trim_rows(rows))])


def _html(data: bytes) -> Extracted:
    tree = parse_html(_decode(data))
    for bad in tree.css("script,style,noscript"):
        bad.decompose()
    title_node = tree.css_first("title")
    title = node_text(title_node) if title_node else None
    tables: list[Table] = []
    for i, tbl in enumerate(tree.css("table")[:20]):
        rows = [[node_text(c) for c in tr.css("th,td")] for tr in tbl.css("tr")]
        tables.append(Table(f"table{i + 1}", _trim_rows(rows)))
    body = tree.body
    raw = body.text(separator="\n", strip=True) if body is not None else ""
    lines = [collapse(line) for line in raw.splitlines()]
    text = clean_text("\n".join(line for line in lines if line), MAX_TEXT_CHARS)
    return Extracted("html", text, tables, title)


def _text(data: bytes) -> Extracted:
    return Extracted("text", clean_text(_decode(data), MAX_TEXT_CHARS))


def extract(data: bytes, fmt: DocFormat) -> Extracted:
    if fmt == "pdf":
        return _pdf(data)
    if fmt == "docx":
        return _docx(data)
    if fmt == "xlsx":
        return _xlsx(data)
    if fmt == "csv":
        return _csv(data)
    if fmt == "html":
        return _html(data)
    return _text(data)


# --- chunks ---------------------------------------------------------------------------------


def table_lines(table: Table) -> list[str]:
    """Rows as `header: value; header: value` so a chunk is self-explanatory in full-text search."""
    if not table.rows:
        return []
    header, *body = table.rows
    lines: list[str] = []
    for row in body:
        pairs = [
            f"{h}: {v}" if h else v
            for h, v in zip(header + [""] * (len(row) - len(header)), row, strict=False)
            if v
        ]
        if pairs:
            lines.append("; ".join(pairs))
    return lines


def chunk_text(text: str, limit: int = CHUNK_CHARS) -> list[str]:
    """Split on paragraph/line boundaries into pieces of at most `limit` characters."""
    pieces: list[str] = []
    current = ""
    for para in re.split(r"\n{2,}|\n", text):
        para = para.strip()
        if not para:
            continue
        while len(para) > limit:
            cut = para.rfind(" ", 0, limit)
            cut = cut if cut > limit // 2 else limit
            head, para = para[:cut].strip(), para[cut:].strip()
            if current:
                pieces.append(current)
                current = ""
            pieces.append(head)
        if len(current) + len(para) + 1 > limit and current:
            pieces.append(current)
            current = para
        else:
            current = f"{current}\n{para}" if current else para
    if current:
        pieces.append(current)
    return pieces


def plain_text(extracted: Extracted) -> str:
    """The whole readable content: the text, then every table as `header: value` lines."""
    parts: list[str] = []
    if extracted.text:
        parts.append(extracted.text)
    for table in extracted.tables:
        lines = table_lines(table)
        if lines:
            header = f"Таблица {table.name}" if len(extracted.tables) > 1 else ""
            parts.append("\n".join(([header] if header else []) + lines))
    return "\n".join(parts)


def to_chunks(extracted: Extracted) -> list[str]:
    return chunk_text(plain_text(extracted))
